from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class WideColumn:
    index: int
    label: str
    group: str = ''

    @property
    def display(self) -> str:
        if self.group:
            return f'{self.group} / {self.label}'
        return self.label or f'Колонка {self.index + 1}'


@dataclass
class WideTable:
    sheet_name: str
    header_row: int  # 1-based
    columns: list[WideColumn]
    rows: list[dict[int, object]]  # col_index -> value
    preview_rows: list[dict[int, object]]


def list_sheet_names(path: Path | str) -> list[str]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == '.csv':
        return ['CSV']
    if suffix not in {'.xlsx', '.xlsm'}:
        raise ValueError(f'Неподдерживаемый формат: {suffix}')
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _cell_has_value(value) -> bool:
    if value is None:
        return False
    return str(value).replace('\u00a0', ' ').strip() != ''


def _matrix_used_bounds(matrix: list[list]) -> tuple[int, int]:
    """Последние используемые индексы строки/колонки (включительно)."""
    last_row = -1
    last_col = -1
    for row_index, row in enumerate(matrix):
        for col_index, cell in enumerate(row):
            if _cell_has_value(cell):
                last_row = row_index
                last_col = max(last_col, col_index)
    return last_row, last_col


def _trim_matrix_to_used_bounds(matrix: list[list]) -> list[list]:
    """Обрезает хвост пустых строк и колонок."""
    last_row, last_col = _matrix_used_bounds(matrix)
    if last_row < 0:
        return []
    trimmed: list[list] = []
    for row in matrix[: last_row + 1]:
        cells = list(row)
        if len(cells) <= last_col:
            cells.extend([None] * (last_col + 1 - len(cells)))
        else:
            cells = cells[: last_col + 1]
        trimmed.append(cells)
    return trimmed


def _cell_preview(value, *, max_len: int = 18) -> str:
    if value is None:
        return ''
    text = str(value).replace('\u00a0', ' ').strip()
    if len(text) <= max_len:
        return text
    return text[: max_len - 1] + '…'


def build_import_layout_schema(
    path: Path | str,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
    group_row: int = 0,
    max_cols: int | None = None,
    max_data_rows: int | None = None,
) -> dict:
    """
    Схема шапки Excel/CSV для шага настройки импорта.
    Возвращает rows с role: group | header | data | other.
    """
    file_path = Path(path)
    header_row = max(1, int(header_row or 1))
    group_row = max(0, int(group_row or 0))
    if max_data_rows is not None and max_data_rows > 0:
        peek_rows = max(header_row + max_data_rows, group_row + max_data_rows, 8)
    else:
        peek_rows = None
    matrix = _peek_sheet_matrix(file_path, sheet_name=sheet_name, max_rows=peek_rows)
    if not matrix:
        return {
            'header_row': header_row,
            'group_row': group_row,
            'has_groups': group_row > 0,
            'cols_shown': 0,
            'cols_total': 0,
            'data_rows_shown': 0,
            'rows': [],
        }

    width = max((len(row) for row in matrix), default=0)
    if max_cols is not None and max_cols > 0:
        cols_shown = min(max_cols, width) if width else 0
    else:
        cols_shown = width
    if max_data_rows is not None and max_data_rows > 0:
        end_row = min(len(matrix), header_row + max_data_rows)
    else:
        end_row = len(matrix)
    # Показываем с 1-й строки, чтобы были видны группы выше заголовка.
    rows_out: list[dict] = []
    for index in range(0, end_row):
        number = index + 1
        raw = matrix[index]
        cells = [
            _cell_preview(raw[c] if c < len(raw) else None, max_len=36)
            for c in range(cols_shown)
        ]
        if group_row and number == group_row:
            role, role_label = 'group', 'Группы колонок'
        elif number == header_row:
            role, role_label = 'header', 'Заголовки'
        elif number > header_row:
            role, role_label = 'data', 'Данные'
        else:
            role, role_label = 'other', 'Прочее'
        rows_out.append(
            {
                'number': number,
                'role': role,
                'role_label': role_label,
                'cells': cells,
            }
        )
    return {
        'header_row': header_row,
        'group_row': group_row,
        'has_groups': bool(group_row and group_row != header_row),
        'cols_shown': cols_shown,
        'cols_total': width,
        'data_rows_shown': max(0, end_row - header_row),
        'rows': rows_out,
    }


def _peek_sheet_matrix(
    path: Path,
    *,
    sheet_name: str | None,
    max_rows: int | None,
) -> list[list]:
    suffix = path.suffix.lower()
    if suffix == '.csv':
        with path.open('r', encoding='utf-8-sig', newline='') as handle:
            reader = list(csv.reader(handle))
        rows = [list(row) for row in reader]
        rows = rows if max_rows is None else rows[:max_rows]
        return _trim_matrix_to_used_bounds(rows)
    if suffix not in {'.xlsx', '.xlsm'}:
        return []
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries

    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        if sheet_name and sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]
        read_limit = max_rows
        if read_limit is None:
            dimension = sheet.calculate_dimension()
            if dimension:
                _, _, _, read_limit = range_boundaries(dimension)
            else:
                read_limit = sheet.max_row or 1
        matrix: list[list] = []
        for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=read_limit, values_only=True)):
            matrix.append(list(row))
            if max_rows is not None and index + 1 >= max_rows:
                break
        # Не разворачиваем merged на весь лист — только заполняем в пределах превью.
        _expand_merged_cells_limited(sheet, matrix)
        return _trim_matrix_to_used_bounds(matrix)
    finally:
        workbook.close()


def _expand_merged_cells_limited(sheet, matrix: list[list]) -> None:
    """Как _expand_merged_cells, но не раздувает matrix за пределы уже прочитанных строк."""
    if not matrix:
        return
    max_rows = len(matrix)
    ranges = getattr(getattr(sheet, 'merged_cells', None), 'ranges', None) or ()
    for merged in ranges:
        min_row, min_col, max_row, max_col = merged.min_row, merged.min_col, merged.max_row, merged.max_col
        origin_r = min_row - 1
        origin_c = min_col - 1
        if origin_r < 0 or origin_r >= max_rows:
            continue
        origin_row = matrix[origin_r]
        while len(origin_row) <= origin_c:
            origin_row.append(None)
        value = origin_row[origin_c]
        for row_idx in range(origin_r, min(max_row, max_rows)):
            row = matrix[row_idx]
            for col_idx in range(min_col - 1, max_col):
                while len(row) <= col_idx:
                    row.append(None)
                if row_idx == origin_r and col_idx == origin_c:
                    continue
                if row[col_idx] is None or str(row[col_idx]).strip() == '':
                    row[col_idx] = value


def detect_header_layout(
    path: Path | str,
    *,
    sheet_name: str | None = None,
) -> tuple[int, int]:
    """
    Эвристика для сводных: строка групп (редкая) + строка заголовков (плотная).
    Возвращает (header_row, group_row); group_row=0 если групп нет.
    """
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == '.csv':
        return 1, 0
    if suffix not in {'.xlsx', '.xlsm'}:
        return 1, 0

    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]
        rows: list[tuple] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            rows.append(tuple(row))
            if index >= 2:
                break
    finally:
        workbook.close()

    if not rows:
        return 1, 0

    def nonempty(row: tuple) -> int:
        return sum(1 for cell in row if cell is not None and str(cell).strip())

    def joined(row: tuple) -> str:
        return ' '.join(str(cell).casefold() for cell in row if cell is not None and str(cell).strip())

    r1 = nonempty(rows[0])
    r2 = nonempty(rows[1]) if len(rows) > 1 else 0
    text1 = joined(rows[0])
    text2 = joined(rows[1]) if len(rows) > 1 else ''

    identity_tokens = ('наименование', 'название', 'name', 'марка', 'код')
    if any(token in text1 for token in identity_tokens) and r1 >= r2:
        return 1, 0
    if any(token in text2 for token in identity_tokens) and r2 > r1:
        return 2, 1 if r1 > 0 else 0
    if r1 > 0 and r2 >= max(5, r1 * 2):
        return 2, 1
    return 1, 0


def load_wide_table(
    path: Path | str,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
    group_row: int | None = None,
    max_preview: int = 5,
) -> WideTable:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if header_row < 1:
        raise ValueError('Номер строки заголовка должен быть ≥ 1.')
    if suffix == '.csv':
        return _load_csv_wide(file_path, header_row=header_row, max_preview=max_preview)
    if suffix in {'.xlsx', '.xlsm'}:
        return _load_xlsx_wide(
            file_path,
            sheet_name=sheet_name,
            header_row=header_row,
            group_row=group_row,
            max_preview=max_preview,
        )
    raise ValueError(f'Неподдерживаемый формат: {suffix}')


def _load_csv_wide(path: Path, *, header_row: int, max_preview: int) -> WideTable:
    with path.open('r', encoding='utf-8-sig', newline='') as handle:
        reader = list(csv.reader(handle))
    if not reader:
        return WideTable(sheet_name='CSV', header_row=header_row, columns=[], rows=[], preview_rows=[])
    if header_row > len(reader):
        raise ValueError(f'В файле нет строки заголовка №{header_row}.')
    header = reader[header_row - 1]
    columns = [
        WideColumn(index=i, label=str(cell).strip() if cell else f'col_{i + 1}')
        for i, cell in enumerate(header)
        if str(cell or '').strip()
    ]
    data_rows: list[dict[int, object]] = []
    for raw in reader[header_row:]:
        row = {col.index: raw[col.index] if col.index < len(raw) else None for col in columns}
        if any(v is not None and str(v).strip() != '' for v in row.values()):
            data_rows.append(row)
    return WideTable(
        sheet_name='CSV',
        header_row=header_row,
        columns=columns,
        rows=data_rows,
        preview_rows=data_rows[:max_preview],
    )


def _load_xlsx_wide(
    path: Path,
    *,
    sheet_name: str | None,
    header_row: int,
    group_row: int | None,
    max_preview: int,
) -> WideTable:
    from openpyxl import load_workbook

    # read_only=False: нужны merged_cells (в Excel значение только в левой верхней ячейке).
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        if sheet_name and sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]
            sheet_name = sheet.title

        matrix: list[list] = [list(row) for row in sheet.iter_rows(values_only=True)]
        _expand_merged_cells(sheet, matrix)

        if header_row > len(matrix):
            raise ValueError(f'На листе «{sheet_name}» нет строки заголовка №{header_row}.')

        header = matrix[header_row - 1]
        groups = matrix[group_row - 1] if group_row and 1 <= group_row <= len(matrix) else None
        width = max(len(header), len(groups) if groups is not None else 0)
        for raw in matrix[header_row:]:
            width = max(width, len(raw))

        # протащить group labels вправо (merged-like / незакрытые объединения)
        filled_groups: list[str] = []
        if groups is not None:
            last = ''
            for index in range(width):
                cell = groups[index] if index < len(groups) else None
                text = str(cell).strip() if cell is not None else ''
                if text:
                    last = text
                filled_groups.append(last)

        data_matrix = matrix[header_row:]

        def column_has_data(index: int) -> bool:
            for raw in data_matrix:
                if index < len(raw) and raw[index] is not None and str(raw[index]).strip():
                    return True
            return False

        columns: list[WideColumn] = []
        for index in range(width):
            cell = header[index] if index < len(header) else None
            label = str(cell).strip() if cell is not None else ''
            group = filled_groups[index] if index < len(filled_groups) else ''
            if not label:
                # Не отбрасываем колонки с пустым заголовком (частая ошибка: строка групп как header)
                if not group and not column_has_data(index):
                    continue
                label = group or f'Колонка {index + 1}'
                group = group if group != label else ''
            columns.append(WideColumn(index=index, label=label, group=group))

        data_rows: list[dict[int, object]] = []
        for raw in data_matrix:
            row = {
                col.index: raw[col.index] if col.index < len(raw) else None
                for col in columns
            }
            if not any(v is not None and str(v).strip() != '' for v in row.values()):
                continue
            if _is_header_echo_row(row, columns):
                continue
            data_rows.append(row)

        return WideTable(
            sheet_name=sheet_name or sheet.title,
            header_row=header_row,
            columns=columns,
            rows=data_rows,
            preview_rows=data_rows[:max_preview],
        )
    finally:
        workbook.close()


def _expand_merged_cells(sheet, matrix: list[list]) -> None:
    """Копирует значение из левой верхней ячейки объединения во все ячейки диапазона."""
    ranges = getattr(getattr(sheet, 'merged_cells', None), 'ranges', None) or ()
    for merged in ranges:
        min_row = merged.min_row
        max_row = merged.max_row
        min_col = merged.min_col
        max_col = merged.max_col
        origin_r = min_row - 1
        origin_c = min_col - 1
        while len(matrix) <= origin_r:
            matrix.append([])
        origin_row = matrix[origin_r]
        while len(origin_row) <= origin_c:
            origin_row.append(None)
        value = origin_row[origin_c]
        for row_idx in range(min_row - 1, max_row):
            while len(matrix) <= row_idx:
                matrix.append([])
            row = matrix[row_idx]
            for col_idx in range(min_col - 1, max_col):
                while len(row) <= col_idx:
                    row.append(None)
                if row_idx == origin_r and col_idx == origin_c:
                    continue
                row[col_idx] = value


def _is_header_echo_row(row: dict[int, object], columns: list[WideColumn]) -> bool:
    """True, если строка повторяет заголовки (частый эффект неверной строки header)."""
    values = []
    label_matches = 0
    for col in columns:
        raw = row.get(col.index)
        if raw is None or not str(raw).strip():
            continue
        text = str(raw).replace('\u00a0', ' ').strip().casefold()
        values.append(text)
        label = (col.label or '').strip().casefold()
        if text == label:
            label_matches += 1

    if not values:
        return False
    # Классика сводной при header_row=1: первая клетка = «Наименование»
    if values[0] in _HEADER_LIKE_VALUES:
        return True
    if label_matches >= 3:
        return True
    title_like = sum(1 for value in values[:8] if _looks_like_column_title(value))
    if title_like >= 4:
        return True
    return False


def _looks_like_column_title(text: str) -> bool:
    if text in _HEADER_LIKE_VALUES:
        return True
    return bool(
        re.search(
            r'(нагрузка|прочность|модуль|удлинение|плотность|толщина|плетен|замаслив|производител)',
            text,
            re.IGNORECASE,
        )
    )


_HEADER_LIKE_VALUES = frozenset({
    'наименование',
    'название',
    'name',
    'марка',
    'код',
    'code',
    'артикул',
    'тип плетения',
    'производитель',
    'описание',
    'плотность пов',
    'замасливатель',
})
