from __future__ import annotations

import csv
from pathlib import Path

from apps.materials.imports.schema import is_blank_row, normalize_row


def read_import_file(path: Path | str) -> list[dict]:
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f'Файл не найден: {file_path}')
    suffix = file_path.suffix.lower()
    if suffix == '.csv':
        return read_csv(file_path)
    if suffix in {'.xlsx', '.xlsm'}:
        return read_xlsx(file_path)
    raise ValueError(f'Неподдерживаемый формат файла: {suffix or "(без расширения)"}')


def read_csv(path: Path) -> list[dict]:
    rows: list[dict] = []
    with path.open('r', encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            return rows
        for index, raw_row in enumerate(reader, start=2):
            row = normalize_row(raw_row, row_number=index)
            if is_blank_row(row):
                continue
            rows.append(row)
    return rows


def read_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    from apps.materials.imports.wide import _expand_merged_cells

    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = workbook.active
        matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
        _expand_merged_cells(sheet, matrix)
        if not matrix:
            return []
        headers = [str(cell).strip() if cell is not None else '' for cell in matrix[0]]
        rows: list[dict] = []
        for offset, values in enumerate(matrix[1:], start=2):
            raw_row = {
                headers[index]: values[index] if index < len(values) else None
                for index in range(len(headers))
                if headers[index]
            }
            row = normalize_row(raw_row, row_number=offset)
            if is_blank_row(row):
                continue
            rows.append(row)
        return rows
    finally:
        workbook.close()
